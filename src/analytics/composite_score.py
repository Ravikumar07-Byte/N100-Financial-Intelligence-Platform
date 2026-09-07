"""
N100 Financial Intelligence Platform
Sprint 3 - Day 17
Composite Quality Score and Excel Export

Implements:
- 0-100 composite quality score
- P10/P90 winsorisation
- Sector-relative normalization
- Profitability score
- Cash quality score
- Growth score
- Leverage score
- Preset-level Excel export
- Threshold-based green/red formatting
"""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


ROOT_DIR = Path(__file__).resolve().parents[2]

OUTPUT_DIR = ROOT_DIR / "output"
OUTPUT_FILE = OUTPUT_DIR / "screener_output.xlsx"


# =====================================================================
# SCORE WEIGHTS
# =====================================================================

WEIGHTS = {
    # Profitability = 35%
    "roe": 15,
    "roce": 10,
    "npm": 10,

    # Cash Quality = 30%
    "fcf_cagr": 15,
    "cfo_pat": 10,
    "fcf_positive": 5,

    # Growth = 20%
    "revenue_cagr": 10,
    "pat_cagr": 10,

    # Leverage = 15%
    "debt_to_equity": 10,
    "icr": 5,
}


# =====================================================================
# HELPER FUNCTIONS
# =====================================================================

def winsorize_and_scale(
    series: pd.Series,
    higher_is_better: bool = True,
) -> pd.Series:
    """
    Winsorize a metric at P10/P90 and scale it to 0-100.

    Higher-is-better:
        P10 -> 0
        P90 -> 100

    Lower-is-better:
        P10 -> 100
        P90 -> 0

    Missing values receive a neutral score of 50.
    """

    values = pd.to_numeric(
        series,
        errors="coerce",
    )

    valid = values.dropna()

    if valid.empty:
        return pd.Series(
            50.0,
            index=series.index,
            dtype=float,
        )

    p10 = valid.quantile(0.10)
    p90 = valid.quantile(0.90)

    clipped = values.clip(
        lower=p10,
        upper=p90,
    )

    if p90 == p10:
        score = pd.Series(
            50.0,
            index=series.index,
            dtype=float,
        )
    else:
        score = (
            (clipped - p10)
            / (p90 - p10)
            * 100
        )

    if not higher_is_better:
        score = 100 - score

    return score.fillna(50.0)


# =====================================================================
# COMPOSITE SCORE CALCULATOR
# =====================================================================

class CompositeScoreCalculator:
    """
    Calculate Sprint 3 Day 17 composite quality scores.
    """

    def __init__(self):

        self.required_columns = [
            "company_id",
            "broad_sector",
            "return_on_equity_pct",
            "return_on_capital_employed_pct",
            "net_profit_margin_pct",
            "free_cash_flow_cr",
            "cash_from_operations_cr",
            "net_profit",
            "revenue_cagr_5yr",
            "pat_cagr_5yr",
            "debt_to_equity",
            "effective_icr",
        ]

    # -----------------------------------------------------------------
    # Validation
    # -----------------------------------------------------------------

    def validate_columns(
        self,
        df: pd.DataFrame,
    ) -> None:
        """Validate required Day 17 columns."""

        missing = [
            column
            for column in self.required_columns
            if column not in df.columns
        ]

        if missing:
            raise KeyError(
                "Missing required Day 17 columns: "
                f"{missing}"
            )

    # -----------------------------------------------------------------
    # Historical FCF CAGR
    # -----------------------------------------------------------------

    def calculate_fcf_cagr(
        self,
        historical_df: pd.DataFrame,
    ) -> pd.DataFrame:
        """
        Calculate FCF CAGR over a true 5-year period.

        Requires:
            company_id
            year
            free_cash_flow_cr

        A value is calculated only when:
            - starting observation is approximately 5 years before
              the latest observation
            - starting FCF > 0
            - ending FCF > 0

        Otherwise the result is NaN.

        This avoids incorrectly calling a 3-year CAGR a 5-year CAGR.
        """

        columns = [
            "company_id",
            "fcf_cagr_5yr",
        ]

        if historical_df is None or historical_df.empty:
            return pd.DataFrame(columns=columns)

        required = {
            "company_id",
            "year",
            "free_cash_flow_cr",
        }

        missing = required - set(
            historical_df.columns
        )

        if missing:
            raise KeyError(
                "Missing FCF history columns: "
                f"{sorted(missing)}"
            )

        results = []

        history = historical_df.copy()

        history["year_date"] = pd.to_datetime(
            history["year"],
            format="%Y-%m",
            errors="coerce",
        )

        history["free_cash_flow_cr"] = pd.to_numeric(
            history["free_cash_flow_cr"],
            errors="coerce",
        )

        for company_id, group in history.groupby(
            "company_id"
        ):

            group = (
                group
                .dropna(subset=["year_date"])
                .sort_values("year_date")
            )

            group = group[
                group["free_cash_flow_cr"].notna()
            ]

            fcf_cagr = np.nan

            if len(group) >= 2:

                latest_row = group.iloc[-1]

                latest_date = latest_row["year_date"]
                latest_fcf = latest_row[
                    "free_cash_flow_cr"
                ]

                target_date = (
                    latest_date
                    - pd.DateOffset(years=5)
                )

                # Find the historical observation
                # closest to exactly five years ago.
                group["date_difference"] = (
                    (
                        group["year_date"]
                        - target_date
                    )
                    .abs()
                    .dt.days
                )

                start_row = group.loc[
                    group["date_difference"].idxmin()
                ]

                start_date = start_row["year_date"]
                start_fcf = start_row[
                    "free_cash_flow_cr"
                ]

                actual_years = (
                    latest_date - start_date
                ).days / 365.25

                if (
                    pd.notna(start_fcf)
                    and pd.notna(latest_fcf)
                    and start_fcf > 0
                    and latest_fcf > 0
                    and actual_years >= 4.5
                ):
                    fcf_cagr = (
                        (
                            latest_fcf
                            / start_fcf
                        )
                        ** (1 / actual_years)
                        - 1
                    ) * 100

            results.append(
                {
                    "company_id": company_id,
                    "fcf_cagr_5yr": fcf_cagr,
                }
            )

        return pd.DataFrame(results)

    # -----------------------------------------------------------------
    # CFO / PAT
    # -----------------------------------------------------------------

    @staticmethod
    def calculate_cfo_pat_ratio(
        df: pd.DataFrame,
    ) -> pd.Series:
        """
        Calculate CFO/PAT ratio.
        """

        cfo = pd.to_numeric(
            df["cash_from_operations_cr"],
            errors="coerce",
        )

        pat = pd.to_numeric(
            df["net_profit"],
            errors="coerce",
        )

        ratio = pd.Series(
            np.nan,
            index=df.index,
            dtype=float,
        )

        valid = (
            pat.notna()
            & cfo.notna()
            & (pat != 0)
        )

        ratio.loc[valid] = (
            cfo.loc[valid]
            / pat.loc[valid]
        )

        return ratio

    # -----------------------------------------------------------------
    # Sector Relative Normalization
    # -----------------------------------------------------------------

    @staticmethod
    def sector_normalize(
        df: pd.DataFrame,
        column: str,
        higher_is_better: bool = True,
    ) -> pd.Series:
        """
        Normalize a metric independently within each broad sector.
        """

        return (
            df.groupby(
                "broad_sector",
                dropna=False,
                group_keys=False,
            )[column]
            .transform(
                lambda series: winsorize_and_scale(
                    series,
                    higher_is_better=higher_is_better,
                )
            )
        )

    # -----------------------------------------------------------------
    # Main Calculation
    # -----------------------------------------------------------------

    def calculate(
        self,
        df: pd.DataFrame,
        historical_df: pd.DataFrame | None = None,
    ) -> pd.DataFrame:
        """
        Calculate the complete Day 17 composite quality score.

        Returns the original dataframe plus:
        - fcf_cagr_5yr
        - cfo_pat_ratio
        - fcf_positive_flag
        - individual metric scores
        - component scores
        - composite_quality_score
        """

        self.validate_columns(df)

        result = df.copy()

        # -------------------------------------------------------------
        # Ensure numeric fields
        # -------------------------------------------------------------

        numeric_columns = [
            "return_on_equity_pct",
            "return_on_capital_employed_pct",
            "net_profit_margin_pct",
            "free_cash_flow_cr",
            "cash_from_operations_cr",
            "net_profit",
            "revenue_cagr_5yr",
            "pat_cagr_5yr",
            "debt_to_equity",
            "effective_icr",
        ]

        for column in numeric_columns:

            result[column] = pd.to_numeric(
                result[column],
                errors="coerce",
            )

        # -------------------------------------------------------------
        # Historical FCF CAGR
        # -------------------------------------------------------------

        if "fcf_cagr_5yr" not in result.columns:

            if historical_df is not None:

                fcf_cagr = (
                    self.calculate_fcf_cagr(
                        historical_df
                    )
                )

                result = result.merge(
                    fcf_cagr,
                    on="company_id",
                    how="left",
                )

            else:

                result["fcf_cagr_5yr"] = np.nan

        # -------------------------------------------------------------
        # CFO/PAT ratio
        # -------------------------------------------------------------

        result["cfo_pat_ratio"] = (
            self.calculate_cfo_pat_ratio(
                result
            )
        )

        # -------------------------------------------------------------
        # FCF positive flag
        # -------------------------------------------------------------

        result["fcf_positive_flag"] = (
            pd.to_numeric(
                result["free_cash_flow_cr"],
                errors="coerce",
            )
            > 0
        ).astype(int)

        # -------------------------------------------------------------
        # Metric scores
        # -------------------------------------------------------------

        result["score_roe"] = (
            self.sector_normalize(
                result,
                "return_on_equity_pct",
                higher_is_better=True,
            )
        )

        result["score_roce"] = (
            self.sector_normalize(
                result,
                "return_on_capital_employed_pct",
                higher_is_better=True,
            )
        )

        result["score_npm"] = (
            self.sector_normalize(
                result,
                "net_profit_margin_pct",
                higher_is_better=True,
            )
        )

        result["score_fcf_cagr"] = (
            self.sector_normalize(
                result,
                "fcf_cagr_5yr",
                higher_is_better=True,
            )
        )

        result["score_cfo_pat"] = (
            self.sector_normalize(
                result,
                "cfo_pat_ratio",
                higher_is_better=True,
            )
        )

        result["score_fcf_positive"] = (
            result["fcf_positive_flag"] * 100
        )

        result["score_revenue_cagr"] = (
            self.sector_normalize(
                result,
                "revenue_cagr_5yr",
                higher_is_better=True,
            )
        )

        result["score_pat_cagr"] = (
            self.sector_normalize(
                result,
                "pat_cagr_5yr",
                higher_is_better=True,
            )
        )

        result["score_debt_to_equity"] = (
            self.sector_normalize(
                result,
                "debt_to_equity",
                higher_is_better=False,
            )
        )

        result["score_icr"] = (
            self.sector_normalize(
                result,
                "effective_icr",
                higher_is_better=True,
            )
        )

        # -------------------------------------------------------------
        # Weighted component scores
        # -------------------------------------------------------------

        # Profitability = 35%
        result["profitability_score"] = (
            result["score_roe"] * 0.15
            + result["score_roce"] * 0.10
            + result["score_npm"] * 0.10
        )

        # Cash Quality = 30%
        result["cash_quality_score"] = (
            result["score_fcf_cagr"] * 0.15
            + result["score_cfo_pat"] * 0.10
            + result["score_fcf_positive"] * 0.05
        )

        # Growth = 20%
        result["growth_score"] = (
            result["score_revenue_cagr"] * 0.10
            + result["score_pat_cagr"] * 0.10
        )

        # Leverage = 15%
        result["leverage_score"] = (
            result["score_debt_to_equity"] * 0.10
            + result["score_icr"] * 0.05
        )

        # -------------------------------------------------------------
        # Final composite score
        # -------------------------------------------------------------

        result["composite_quality_score"] = (
            result["profitability_score"]
            + result["cash_quality_score"]
            + result["growth_score"]
            + result["leverage_score"]
        )

        result["composite_quality_score"] = (
            result["composite_quality_score"]
            .clip(0, 100)
            .round(2)
        )

        return result


# =====================================================================
# EXCEL EXPORTER
# =====================================================================

class ScreenerExcelExporter:
    """
    Generate Day 17 screener Excel workbook.
    """

    GREEN = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    RED = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    HEADER = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    def __init__(
        self,
        output_file: Path = OUTPUT_FILE,
    ):
        self.output_file = Path(
            output_file
        )

        self.output_file.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

    # -----------------------------------------------------------------
    # KPI Selection
    # -----------------------------------------------------------------

    @staticmethod
    def select_kpis(
        df: pd.DataFrame,
    ) -> pd.DataFrame:
        """
        Select exactly 20 KPI/output columns.
        """

        kpis = [
            "company_id",
            "company_name",
            "broad_sector",
            "sub_sector",

            "return_on_equity_pct",
            "return_on_capital_employed_pct",
            "net_profit_margin_pct",

            "free_cash_flow_cr",
            "fcf_cagr_5yr",
            "cash_from_operations_cr",
            "cfo_pat_ratio",

            "revenue_cagr_5yr",
            "pat_cagr_5yr",

            "debt_to_equity",
            "effective_icr",

            "pe_ratio",
            "pb_ratio",
            "dividend_yield_pct",
            "dividend_payout_ratio_pct",

            "composite_quality_score",
        ]

        # Exactly 20 columns
        assert len(kpis) == 20

        result = df.copy()

        for column in kpis:

            if column not in result.columns:
                result[column] = np.nan

        return result[kpis].copy()

    # -----------------------------------------------------------------
    # Threshold Evaluation
    # -----------------------------------------------------------------

    @staticmethod
    def _cell_passes_threshold(
        row: pd.Series,
        key: str,
        threshold,
    ) -> bool:
        """
        Check whether a single KPI passes its threshold.
        """

        column_map = {
            "roe_min": "return_on_equity_pct",
            "free_cash_flow_min": "free_cash_flow_cr",
            "revenue_cagr_5yr_min": "revenue_cagr_5yr",
            "pat_cagr_5yr_min": "pat_cagr_5yr",
            "pe_max": "pe_ratio",
            "pb_max": "pb_ratio",
            "dividend_yield_min": "dividend_yield_pct",
            "dividend_payout_ratio_max":
                "dividend_payout_ratio_pct",
            "sales_min": "sales",
            "revenue_cagr_3yr_min":
                "revenue_cagr_3yr",
        }

        # -------------------------------------------------------------
        # Debt-to-equity
        # -------------------------------------------------------------

        if key == "debt_to_equity_max":

            sector = str(
                row.get(
                    "broad_sector",
                    "",
                )
            )

            # Financial companies are exempt from the
            # normal D/E screening rule.
            if sector == "Financials":
                return True

            value = pd.to_numeric(
                row.get("debt_to_equity"),
                errors="coerce",
            )

            if pd.isna(value):
                return False

            if float(threshold) == 0:
                return value == 0

            return value < float(threshold)

        # -------------------------------------------------------------
        # Generic metric
        # -------------------------------------------------------------

        if key not in column_map:
            return True

        column = column_map[key]

        value = pd.to_numeric(
            row.get(column),
            errors="coerce",
        )

        if pd.isna(value):
            return False

        threshold = float(threshold)

        if key.endswith("_min"):
            return value > threshold

        if key.endswith("_max"):
            return value < threshold

        return True

    # -----------------------------------------------------------------
    # Workbook
    # -----------------------------------------------------------------

    def export(
        self,
        screener_results: dict[str, pd.DataFrame],
        screener_configs: dict[str, dict],
    ) -> Path:
        """
        Generate one Excel sheet per screener preset.
        """

        with pd.ExcelWriter(
            self.output_file,
            engine="openpyxl",
        ) as writer:

            for screener_name, df in (
                screener_results.items()
            ):

                result = df.copy()

                result = result.sort_values(
                    "composite_quality_score",
                    ascending=False,
                    na_position="last",
                )

                result = self.select_kpis(
                    result
                )

                sheet_name = screener_name[:31]

                result.to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False,
                )

        self._format_workbook(
            screener_results,
            screener_configs,
        )

        return self.output_file

    # -----------------------------------------------------------------
    # Formatting
    # -----------------------------------------------------------------

    def _format_workbook(
        self,
        screener_results: dict[str, pd.DataFrame],
        screener_configs: dict[str, dict],
    ) -> None:
        """
        Apply workbook formatting.

        Threshold cells:
            Green = passes threshold
            Red   = fails threshold

        Unlike the previous implementation, only the
        relevant KPI cell is coloured.
        """

        workbook = load_workbook(
            self.output_file
        )

        for screener_name, original_df in (
            screener_results.items()
        ):

            sheet_name = screener_name[:31]

            if sheet_name not in workbook.sheetnames:
                continue

            ws = workbook[sheet_name]

            # ---------------------------------------------------------
            # Header
            # ---------------------------------------------------------

            for cell in ws[1]:
                cell.fill = self.HEADER

            ws.freeze_panes = "A2"

            # ---------------------------------------------------------
            # Column widths
            # ---------------------------------------------------------

            for column_cells in ws.columns:

                max_length = 0

                for cell in column_cells:

                    value = (
                        ""
                        if cell.value is None
                        else str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        len(value),
                    )

                column_letter = (
                    column_cells[0]
                    .column_letter
                )

                ws.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    35,
                )

            # ---------------------------------------------------------
            # Filter configuration
            # ---------------------------------------------------------

            filters = (
                screener_configs[
                    screener_name
                ].get(
                    "filters",
                    {},
                )
            )

            # ---------------------------------------------------------
            # Sort original dataframe exactly like export
            # ---------------------------------------------------------

            sorted_original = (
                original_df.copy()
                .sort_values(
                    "composite_quality_score",
                    ascending=False,
                    na_position="last",
                )
                .reset_index(drop=True)
            )

            # ---------------------------------------------------------
            # Header -> column lookup
            # ---------------------------------------------------------

            headers = {
                cell.value: cell.column
                for cell in ws[1]
            }

            # ---------------------------------------------------------
            # Threshold cell formatting
            # ---------------------------------------------------------

            for row_index in range(
                2,
                ws.max_row + 1,
            ):

                source_index = row_index - 2

                if (
                    source_index
                    >= len(sorted_original)
                ):
                    continue

                source_row = sorted_original.iloc[
                    source_index
                ]

                for key, threshold in filters.items():

                    # Find corresponding KPI column
                    column_map = {
                        "roe_min":
                            "return_on_equity_pct",

                        "free_cash_flow_min":
                            "free_cash_flow_cr",

                        "revenue_cagr_5yr_min":
                            "revenue_cagr_5yr",

                        "pat_cagr_5yr_min":
                            "pat_cagr_5yr",

                        "pe_max":
                            "pe_ratio",

                        "pb_max":
                            "pb_ratio",

                        "dividend_yield_min":
                            "dividend_yield_pct",

                        "dividend_payout_ratio_max":
                            "dividend_payout_ratio_pct",

                        "sales_min":
                            "sales",

                        "revenue_cagr_3yr_min":
                            "revenue_cagr_3yr",

                        "debt_to_equity_max":
                            "debt_to_equity",
                    }

                    column_name = column_map.get(
                        key
                    )

                    if column_name is None:
                        continue

                    if column_name not in headers:
                        continue

                    passes = (
                        self._cell_passes_threshold(
                            source_row,
                            key,
                            threshold,
                        )
                    )

                    excel_column = headers[
                        column_name
                    ]

                    cell = ws.cell(
                        row=row_index,
                        column=excel_column,
                    )

                    cell.fill = (
                        self.GREEN
                        if passes
                        else self.RED
                    )

            # ---------------------------------------------------------
            # Number formats
            # ---------------------------------------------------------

            percentage_columns = {
                "return_on_equity_pct",
                "return_on_capital_employed_pct",
                "net_profit_margin_pct",
                "fcf_cagr_5yr",
                "cfo_pat_ratio",
                "revenue_cagr_5yr",
                "pat_cagr_5yr",
                "dividend_yield_pct",
                "dividend_payout_ratio_pct",
                "composite_quality_score",
            }

            decimal_columns = {
                "debt_to_equity",
                "pe_ratio",
                "pb_ratio",
                "effective_icr",
            }

            for row in range(
                2,
                ws.max_row + 1,
            ):

                for header, column in headers.items():

                    cell = ws.cell(
                        row=row,
                        column=column,
                    )

                    if header in percentage_columns:

                        cell.number_format = (
                            "0.00"
                        )

                    elif header in decimal_columns:

                        cell.number_format = (
                            "0.00"
                        )

        workbook.save(
            self.output_file
        )


# =====================================================================
# CONVENIENCE FUNCTION
# =====================================================================

def generate_screener_export(
    screener_results: dict[str, pd.DataFrame],
    screener_configs: dict[str, dict],
) -> Path:
    """
    Generate the Day 17 Excel export.
    """

    exporter = ScreenerExcelExporter()

    return exporter.export(
        screener_results,
        screener_configs,
    )